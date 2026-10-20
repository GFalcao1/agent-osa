from __future__ import annotations

import datetime as dt
import io
import re
import zipfile
from pathlib import Path
from types import ModuleType

import pytest
from PIL import Image

from osa.core.settings import Settings
from osa.domain.errors import ExtractionFailed
from osa.extract import ExtractedContent, ExtractLimits, extract

PNG_MAGIC = bytes.fromhex("89504e470d0a1a0a")


def digits(text: str) -> str:
    return re.sub(r"\D", "", text)


def assert_pngs(content: ExtractedContent) -> None:
    for image in content.images:
        assert image.data.startswith(PNG_MAGIC)
        with Image.open(io.BytesIO(image.data)) as decoded:
            assert decoded.size == (image.width, image.height)
            assert max(decoded.size) <= 1600


def assert_nota_text(text: str, fixtures: ModuleType) -> None:
    for word in fixtures.NOTA_WORDS:
        assert word in text, word
    assert fixtures.CNPJ_NOTA in digits(text)
    assert fixtures.VALOR_NOTA in text.replace("R$ ", "R$")


# --- PDF -------------------------------------------------------------------


def test_textual_pdf_returns_the_text_layer(
    docs: dict[str, Path], fixture_maker: ModuleType
) -> None:
    content = extract(docs["texto.pdf"])

    assert content.method == "pdf_text"
    assert content.images == ()
    assert content.pages_total == 1
    assert content.truncated is False
    for line in fixture_maker.NOTA_LINES:
        assert line in content.text
    assert fixture_maker.CNPJ_NOTA_MASKED in content.text


def test_scanned_pdf_ocr_returns_the_fixture_words(
    docs: dict[str, Path], fixture_maker: ModuleType
) -> None:
    content = extract(docs["escaneado.pdf"])

    assert content.method == "pdf_scanned"
    assert content.pages_total == 1
    assert len(content.images) == 1
    assert content.images[0].page == 1
    assert_pngs(content)
    assert content.text != ""  # a scan is never reported as empty
    assert_nota_text(content.text, fixture_maker)


def test_scanned_pdf_pages_beyond_the_image_limit_are_still_read_by_ocr(
    docs: dict[str, Path],
) -> None:
    content = extract(docs["escaneado_7p.pdf"], ExtractLimits(max_image_pages=5))

    assert content.method == "pdf_scanned"
    assert content.pages_total == 7
    assert [image.page for image in content.images] == [1, 2, 3, 4, 5]
    assert content.truncated is False
    for page in range(1, 8):
        assert f"Página {page} de 7" in content.text, page


def test_ocr_budget_marks_the_pages_left_out(docs: dict[str, Path]) -> None:
    content = extract(docs["escaneado_7p.pdf"], ExtractLimits(max_image_pages=2, max_ocr_pages=3))

    assert len(content.images) == 2
    assert "Página 3 de 7" in content.text
    assert "Página 4 de 7" not in content.text
    assert content.truncated is True


def test_mixed_pdf_is_decided_page_by_page(
    docs: dict[str, Path], fixture_maker: ModuleType
) -> None:
    content = extract(docs["misto.pdf"])

    assert content.method == "pdf_mixed"
    assert content.pages_total == 3
    assert "Página 1 com camada de texto" in content.text
    assert "Página 3 com camada de texto" in content.text
    # The image-only page is not lost: it is an image and its OCR text is present.
    assert [image.page for image in content.images] == [2]
    assert "somente imagem" in content.text
    assert fixture_maker.CNPJ_CONTRATO in digits(content.text)
    assert content.truncated is False


def test_scanned_page_with_a_scanner_stamp_is_still_an_image_page(
    docs: dict[str, Path], fixture_maker: ModuleType
) -> None:
    # The stamp alone passes min_chars_per_page; the full-page image decides.
    content = extract(docs["escaneado_carimbo.pdf"])

    assert content.method == "pdf_scanned"
    assert [image.page for image in content.images] == [1]
    assert_pngs(content)
    assert "lote 000123" in content.text  # the text layer is kept
    assert_nota_text(content.text, fixture_maker)  # and the scan went through OCR


LETTER_PAGES = [
    [
        f"Carta fictícia, página {n} de 4, em papel timbrado digitalizado como fundo.",
        f"Parágrafo {n}: a Exemplo Comércio e Serviços Ltda confirma o recebimento.",
        f"Parágrafo {n}: o valor de R$ 12.345,67 será pago até 15/11/2026 sem multa.",
        f"Parágrafo {n}: dúvidas pelo setor financeiro fictício, ramal 000{n}.",
        f"Encerramento da página {n} da carta fictícia com camada de texto real.",
    ]
    for n in range(1, 5)
]


def letterhead_pdf(path: Path) -> Path:
    """Four pages of real text drawn over a full-page background image (letterhead)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    background = Image.new("RGB", (620, 877), (250, 250, 245))
    background.paste((20, 40, 120), (40, 30, 580, 90))  # the letterhead band
    buffer = io.BytesIO()
    background.save(buffer, format="PNG")
    pdf = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    for lines in LETTER_PAGES:
        pdf.drawImage(ImageReader(io.BytesIO(buffer.getvalue())), 0, 0, A4[0], A4[1])
        pdf.setFont("Helvetica", 11)
        for row, line in enumerate(lines):
            pdf.drawString(56, 700 - row * 24, line)
        pdf.showPage()
    pdf.save()
    return path


def test_letterhead_letter_keeps_its_text_layer_without_ocr(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from osa.extract import pdf as pdf_module

    path = letterhead_pdf(tmp_path / "timbrado.pdf")

    def no_ocr(*args: object, **kwargs: object) -> str:
        pytest.fail("a page with a large text layer went through OCR")

    monkeypatch.setattr(pdf_module, "ocr_image", no_ocr)
    content = pdf_module.extract_pdf(path, ExtractLimits())

    assert content.method == "pdf_text"
    assert content.pages_total == 4
    assert content.truncated is False
    for lines in LETTER_PAGES:
        for line in lines:
            assert content.text.count(line) == 1, line
    # Still images for the vision model (the letterhead may matter), within the limit.
    assert [image.page for image in content.images] == [1, 2, 3, 4]
    assert_pngs(content)
    limited = pdf_module.extract_pdf(path, ExtractLimits(max_image_pages=2, max_ocr_pages=2))
    assert [image.page for image in limited.images] == [1, 2]
    assert limited.text == content.text


def test_letterhead_through_the_isolated_worker(tmp_path: Path) -> None:
    content = extract(letterhead_pdf(tmp_path / "timbrado.pdf"))

    assert (content.method, content.pages_total, content.truncated) == ("pdf_text", 4, False)
    assert content.text.count("Carta fictícia, página") == 4


def test_text_over_image_threshold_is_a_validated_limit() -> None:
    assert ExtractLimits().min_chars_text_over_image == 10 * ExtractLimits().min_chars_per_page
    with pytest.raises(ValueError, match="min_chars_text_over_image"):
        ExtractLimits(min_chars_per_page=50, min_chars_text_over_image=40)


def test_text_page_with_a_small_logo_stays_a_text_page(tmp_path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    logo = io.BytesIO()
    Image.new("RGB", (120, 60), "navy").save(logo, format="PNG")
    path = tmp_path / "logo.pdf"
    pdf = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    pdf.drawImage(ImageReader(io.BytesIO(logo.getvalue())), 40, 760, width=120, height=60)
    pdf.setFont("Helvetica", 12)
    pdf.drawString(56, 700, "Carta fictícia com logotipo pequeno no cabeçalho.")
    pdf.showPage()
    pdf.save()

    content = extract(path)

    assert content.method == "pdf_text"
    assert content.images == ()


def form_scan_pdf(path: Path, scan: bytes, *, nested: bool, stamp: str | None) -> Path:
    """A page-sized scan drawn at 1/10 inside a Form XObject that the page scales up 10x."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    width, height = A4
    pdf = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    pdf.beginForm("inner")
    pdf.drawImage(ImageReader(io.BytesIO(scan)), 0, 0, width=width / 10, height=height / 10)
    pdf.endForm()
    pdf.saveState()
    if nested:  # 2x on the page, 5x inside the outer form
        pdf.beginForm("outer")
        pdf.saveState()
        pdf.scale(5, 5)
        pdf.doForm("inner")
        pdf.restoreState()
        pdf.endForm()
        pdf.scale(2, 2)
        pdf.doForm("outer")
    else:
        pdf.scale(10, 10)
        pdf.doForm("inner")
    pdf.restoreState()
    if stamp:
        pdf.setFont("Helvetica", 6)
        pdf.drawString(20, 10, stamp)
    pdf.showPage()
    pdf.save()
    return path


def raw_form_matrix_pdf(form_matrix: str, cm: str) -> bytes:
    """Hand-written PDF: a 1x1 image filling a Form whose /Matrix and the page cm scale it."""
    page = f"q {cm} cm /F1 Do Q".encode()
    form = b"q 59.5 0 0 84.2 0 0 cm /Im1 Do Q"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842]"
        b" /Resources << /XObject << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(page), page),
        b"<< /Type /XObject /Subtype /Form /BBox [0 0 59.5 84.2] /Matrix [%s]"
        b" /Resources << /XObject << /Im1 6 0 R >> >> /Length %d >>\nstream\n%s\nendstream"
        % (form_matrix.encode(), len(form), form),
        b"<< /Type /XObject /Subtype /Image /Width 1 /Height 1 /ColorSpace /DeviceGray"
        b" /BitsPerComponent 8 /Length 1 >>\nstream\n\x80\nendstream",
    ]
    out = bytearray(b"%PDF-1.7\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return bytes(out)


@pytest.mark.parametrize("nested", [False, True])
def test_scan_inside_a_scaled_form_is_an_image_page(
    tmp_path: Path, fixture_maker: ModuleType, nested: bool
) -> None:
    scan = fixture_maker._render_png(fixture_maker.NOTA_LINES)
    path = form_scan_pdf(
        tmp_path / "formulario.pdf", scan, nested=nested, stamp=fixture_maker.SCANNER_STAMP
    )

    content = extract(path)

    assert content.method == "pdf_scanned"
    assert [image.page for image in content.images] == [1]
    assert "lote 000123" in content.text
    assert fixture_maker.CNPJ_NOTA in digits(content.text)  # the OCR ran


@pytest.mark.parametrize(
    ("form_matrix", "cm"),
    [
        ("10 0 0 10 0 0", "1 0 0 1 0 0"),
        ("1 0 0 1 0 0", "10 0 0 10 0 0"),
        ("5 0 0 5 0 0", "2 0 0 2 0 0"),
    ],
)
def test_image_coverage_applies_the_form_matrix_and_the_page_transform(
    form_matrix: str, cm: str
) -> None:
    import pypdfium2 as pdfium

    from osa.extract.pdf import image_coverage

    document = pdfium.PdfDocument(raw_form_matrix_pdf(form_matrix, cm))
    try:
        page = document[0]
        assert image_coverage(page) == pytest.approx(1.0, abs=0.01)
        page.close()
    finally:
        document.close()


def test_page_too_complex_to_walk_counts_as_covered(
    docs: dict[str, Path], monkeypatch: pytest.MonkeyPatch
) -> None:
    import pypdfium2 as pdfium

    from osa.extract import pdf as pdf_module

    nested = pdfium.PdfDocument(raw_form_matrix_pdf("1 0 0 1 0 0", "1 0 0 1 0 0"))
    text = pdfium.PdfDocument(str(docs["texto.pdf"]))
    try:
        assert pdf_module.image_coverage(nested[0]) < 0.05  # small image, walk completes
        assert pdf_module.image_coverage(text[0]) == 0.0
        monkeypatch.setattr(pdf_module, "_MAX_FORM_DEPTH", 0)
        assert pdf_module.image_coverage(nested[0]) == 1.0
        monkeypatch.setattr(pdf_module, "_MAX_PAGE_OBJECTS", 1)
        assert pdf_module.image_coverage(text[0]) == 1.0
    finally:
        nested.close()
        text.close()


def test_small_logo_inside_a_form_stays_a_text_page(tmp_path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    logo = io.BytesIO()
    Image.new("RGB", (120, 60), "navy").save(logo, format="PNG")
    path = tmp_path / "logo_form.pdf"
    pdf = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    from reportlab.lib.utils import ImageReader

    pdf.beginForm("logo")
    pdf.drawImage(ImageReader(io.BytesIO(logo.getvalue())), 0, 0, width=12, height=6)
    pdf.endForm()
    pdf.saveState()
    pdf.translate(40, 760)
    pdf.scale(10, 10)
    pdf.doForm("logo")
    pdf.restoreState()
    pdf.setFont("Helvetica", 12)
    pdf.drawString(56, 700, "Carta fictícia com logotipo em formulário reutilizável.")
    pdf.showPage()
    pdf.save()

    content = extract(path)

    assert content.method == "pdf_text"
    assert content.images == ()


def test_password_protected_pdf_fails_cleanly(docs: dict[str, Path]) -> None:
    with pytest.raises(ExtractionFailed):
        extract(docs["protegido.pdf"])


# --- Office ----------------------------------------------------------------


def test_docx_paragraphs_and_tables_in_document_order(
    docs: dict[str, Path], fixture_maker: ModuleType
) -> None:
    content = extract(docs["contrato.docx"])

    assert content.method == "docx"
    assert content.pages_total == 1  # convention: DOCX has no reliable pagination
    assert content.images == ()
    assert content.truncated is False
    text = content.text
    for paragraph in fixture_maker.CONTRATO_PARAGRAPHS:
        assert paragraph in text
    assert "Item | Valor mensal" in text
    assert "Locação de utilitários | R$ 8.900,00" in text
    assert text.count("Observação: reajuste anual") == 1  # merged cell counted once
    assert text.index("Cláusula primeira") < text.index("Item | Valor mensal")
    assert text.index("Item | Valor mensal") < text.index("Assinado em Curitiba")
    # Documented limitation: footers are not extracted.
    assert fixture_maker.CONTRATO_FOOTER not in text


def test_docx_vertically_merged_cell_is_counted_once(tmp_path: Path) -> None:
    from docx import Document

    document = Document()
    table = document.add_table(rows=3, cols=2)
    merged = table.cell(0, 0).merge(table.cell(2, 0))
    merged.text = "MESCLADA-VERTICAL"
    for row in range(3):
        table.cell(row, 1).text = f"B{row + 1}"
    path = tmp_path / "vertical.docx"
    document.save(str(path))

    text = extract(path).text

    assert text.count("MESCLADA-VERTICAL") == 1
    assert text.splitlines() == ["MESCLADA-VERTICAL | B1", " | B2", " | B3"]


def test_xlsx_cells_become_sheet_coordinate_lines(docs: dict[str, Path]) -> None:
    content = extract(docs["planilha.xlsx"])

    assert content.method == "xlsx"
    assert content.pages_total == 2
    assert content.truncated is False
    lines = content.text.splitlines()
    assert "Pagamentos!A1: Fornecedor" in lines
    assert "Pagamentos!A2: Exemplo Peças Ltda" in lines
    assert "Pagamentos!B2: 12345.67" in lines
    assert "Pagamentos!C2: 2026-10-05" in lines
    assert "Notas!A2: Conferido pelo setor financeiro" in lines


def test_xlsx_counts_chartsheets_and_accepts_a_workbook_of_charts_only(tmp_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference

    mixed = Workbook()
    data = mixed.active
    assert data is not None
    data.title = "Dados"
    for value in (1, 2, 3):
        data.append([value])
    chart = BarChart()
    chart.add_data(Reference(data, min_col=1, min_row=1, max_row=3))
    mixed.create_chartsheet("Grafico").add_chart(chart)
    mixed.save(tmp_path / "misto.xlsx")
    only = Workbook()
    only.remove(only.active)
    only.create_chartsheet("Grafico").add_chart(BarChart())
    only.save(tmp_path / "graficos.xlsx")

    with_data = extract(tmp_path / "misto.xlsx")
    charts = extract(tmp_path / "graficos.xlsx")

    assert with_data.pages_total == 2
    assert "Dados!A3: 3" in with_data.text
    assert (charts.method, charts.pages_total, charts.text) == ("xlsx", 1, "")


def make_xlsx(path: Path, rows: int, sheets: int = 1) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    for index in range(sheets):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        assert sheet is not None
        sheet.title = f"S{index + 1}"
        for row in range(1, rows + 1):
            sheet.cell(row=row, column=1, value=f"linha {row}")
    workbook.save(path)
    return path


def test_xlsx_row_and_sheet_limits_mark_truncation(tmp_path: Path) -> None:
    big = make_xlsx(tmp_path / "grande.xlsx", rows=50, sheets=4)

    content = extract(big, ExtractLimits(max_xlsx_rows=10, max_xlsx_sheets=2))

    assert content.pages_total == 4
    assert content.truncated is True
    assert "S1!A10: linha 10" in content.text
    assert "S1!A11" not in content.text
    assert "S2!A1: linha 1" in content.text
    assert "S3!" not in content.text

    exact = extract(make_xlsx(tmp_path / "exato.xlsx", rows=10), ExtractLimits(max_xlsx_rows=10))
    assert exact.truncated is False


# --- Images (OCR in everything) --------------------------------------------


@pytest.mark.parametrize("name", ["imagem.png", "imagem.jpg", "imagem.heic"])
def test_image_ocr_returns_the_fixture_words(
    name: str, docs: dict[str, Path], fixture_maker: ModuleType
) -> None:
    content = extract(docs[name])

    assert content.method == "image"
    assert content.pages_total == 1
    assert len(content.images) == 1
    assert content.images[0].page == 1
    assert_pngs(content)
    assert_nota_text(content.text, fixture_maker)


LINUX_RASTER = Path(__file__).resolve().parents[1] / "fixtures" / "ocr" / "nota_raster_linux.png"


def test_long_line_misread_as_upside_down_is_still_read(fixture_maker: ModuleType) -> None:
    # The nota fixture as pypdfium2 rasterizes it on Linux (CI, ubuntu-24.04): the angle
    # classifier calls the long "Descrição" line upside down (0.94 > cls_thresh 0.9).
    content = extract(LINUX_RASTER)

    assert_nota_text(content.text, fixture_maker)
    assert (
        "Descrição: manutenção preventiva de equipamentos e contratação de consultoria"
        in content.text
    )


def test_upside_down_scan_is_still_read(tmp_path: Path, docs: dict[str, Path]) -> None:
    upside_down = tmp_path / "invertida.png"
    with Image.open(docs["imagem.png"]) as image:
        image.rotate(180).save(upside_down)

    content = extract(upside_down)

    assert "NOTA FISCAL" in content.text
    assert "Prestador" in content.text


def test_recognizer_arbitrates_the_angle_classifier() -> None:
    from osa.extract.ocr import choose_orientation

    originals = ["linha-a", "linha-b", "linha-c"]
    rotated = ["a-rodada", "b-rodada", "linha-c"]
    scores = {"linha-a": 0.99, "a-rodada": 0.40, "linha-b": 0.30, "b-rodada": 0.95}
    calls: list[list[str]] = []

    def recognize(images: list[str]) -> list[float]:
        calls.append(images)
        return [scores[image] for image in images]

    kept = choose_orientation(originals, rotated, [0, 1], recognize)

    # a: the classifier was wrong (upright reads better); b: it was right; c: not flagged.
    assert kept == ["linha-a", "b-rodada", "linha-c"]
    assert calls == [["linha-a", "linha-b", "a-rodada", "b-rodada"]]
    assert choose_orientation(originals, rotated, [], recognize) == rotated
    assert len(calls) == 1  # nothing flagged, no extra recognition


def test_image_without_text_yields_an_image_and_empty_text(docs: dict[str, Path]) -> None:
    content = extract(docs["gradiente.png"])

    assert content.method == "image"
    assert len(content.images) == 1
    assert content.text == ""
    assert content.truncated is False


def test_large_images_are_reduced_for_the_model(tmp_path: Path) -> None:
    big = tmp_path / "grande.png"
    Image.new("RGB", (3000, 2000), "white").save(big)

    content = extract(big, ExtractLimits(max_image_side=800))

    assert (content.images[0].width, content.images[0].height) == (800, 533)


def test_ocr_input_size_matches_the_engine_limit() -> None:
    from osa.extract import ocr

    limits = ExtractLimits()
    params = ocr.engine_params(limits.max_ocr_side)

    # The image is reduced once, by us; the engine must not shrink it again.
    assert params["Global.max_side_len"] == limits.max_ocr_side == 2000
    assert params["EngineConfig.onnxruntime.intra_op_num_threads"] == ocr.OCR_THREADS
    assert params["EngineConfig.onnxruntime.inter_op_num_threads"] == 1


def test_transparent_png_is_composited_over_white(tmp_path: Path) -> None:
    path = tmp_path / "transparente.png"
    Image.new("RGBA", (64, 64), (0, 0, 0, 0)).save(path)

    content = extract(path)

    with Image.open(io.BytesIO(content.images[0].data)) as decoded:
        assert decoded.convert("RGB").getpixel((10, 10)) == (255, 255, 255)


# --- Truncation and limits -------------------------------------------------


def test_text_is_cut_at_the_limit_and_flagged(docs: dict[str, Path]) -> None:
    full = extract(docs["contrato.docx"])

    cut = extract(docs["contrato.docx"], ExtractLimits(max_text_chars=40))
    exact = extract(docs["contrato.docx"], ExtractLimits(max_text_chars=len(full.text)))
    one_short = extract(docs["contrato.docx"], ExtractLimits(max_text_chars=len(full.text) - 1))

    assert len(cut.text) == 40
    assert cut.truncated is True
    assert full.text.startswith(cut.text)
    assert exact.text == full.text
    assert exact.truncated is False
    assert one_short.truncated is True


def test_pdf_text_limit_and_image_limit(docs: dict[str, Path]) -> None:
    cut = extract(docs["texto.pdf"], ExtractLimits(max_text_chars=30))
    assert len(cut.text) == 30
    assert cut.truncated is True

    two = extract(docs["escaneado_7p.pdf"], ExtractLimits(max_image_pages=2, max_ocr_pages=7))
    assert len(two.images) == 2
    assert two.truncated is False


def test_ocr_text_is_cut_at_the_text_limit(docs: dict[str, Path]) -> None:
    content = extract(docs["imagem.png"], ExtractLimits(max_text_chars=25))

    assert len(content.text) == 25
    assert content.truncated is True


def test_limits_are_validated() -> None:
    for name in ("max_text_chars", "max_image_pages", "timeout_s", "max_pixels"):
        with pytest.raises(ValueError, match=name):
            ExtractLimits(**{name: 0})
    with pytest.raises(ValueError, match="max_ocr_pages"):
        ExtractLimits(max_image_pages=5, max_ocr_pages=4)
    with pytest.raises(ValueError, match="max_part_bytes"):
        ExtractLimits(max_part_bytes=2000, max_unzipped_bytes=1000)
    with pytest.raises(ValueError, match="sensitive_extensions"):
        ExtractLimits(sensitive_extensions=())
    with pytest.raises(ValueError, match="sensitive_extensions"):
        ExtractLimits(sensitive_extensions=("pfx",))
    with pytest.raises(ValueError, match="max_text_chars"):
        ExtractLimits(max_text_chars=True)  # type: ignore[arg-type]


def test_limits_from_settings_use_the_existing_settings() -> None:
    settings = Settings(
        max_text_chars=500, max_image_pages=3, sensitive_extensions=(".p7b", ".pfx")
    )

    limits = ExtractLimits.from_settings(settings)

    assert limits.max_text_chars == 500
    assert limits.max_image_pages == 3
    assert limits.sensitive_extensions == (".p7b", ".pfx")
    assert limits.max_ocr_pages >= limits.max_image_pages


def test_limits_from_settings_with_many_image_pages_stay_valid() -> None:
    limits = ExtractLimits.from_settings(Settings(max_image_pages=40))

    assert limits.max_ocr_pages == 40


# --- Errors never escape as anything but ExtractionFailed ------------------


def broken_files(tmp_path: Path, docs: dict[str, Path]) -> dict[str, Path]:
    truncated_zip = tmp_path / "truncado.xlsx"
    truncated_zip.write_bytes(docs["planilha.xlsx"].read_bytes()[:-120])
    cut_png = tmp_path / "cortado.png"
    cut_png.write_bytes(docs["imagem.png"].read_bytes()[:400])
    cut_docx = tmp_path / "cortado.docx"
    cut_docx.write_bytes(docs["contrato.docx"].read_bytes()[:3000])
    garbage_pdf = tmp_path / "lixo.pdf"
    garbage_pdf.write_bytes(b"%PDF-1.7\n" + bytes(range(256)) * 20)
    cut_pdf = tmp_path / "cortado.pdf"
    cut_pdf.write_bytes(docs["texto.pdf"].read_bytes()[:200])
    garbage_jpg = tmp_path / "lixo.jpg"
    garbage_jpg.write_bytes(bytes.fromhex("ffd8ffe0") + bytes(range(256)) * 8)
    bad_zip = tmp_path / "invalido.docx"
    with zipfile.ZipFile(bad_zip, "w") as archive:
        archive.writestr("word/document.xml", "<w:document><w:body>")  # not well formed
    empty = tmp_path / "vazio.pdf"
    empty.write_bytes(b"")
    missing = tmp_path / "sumiu.docx"
    folder = tmp_path / "pasta.jpg"
    folder.mkdir()
    return {
        path.name: path
        for path in (
            truncated_zip,
            cut_png,
            cut_docx,
            garbage_pdf,
            cut_pdf,
            garbage_jpg,
            bad_zip,
            empty,
            missing,
            folder,
        )
    }


def test_corrupt_files_end_as_extraction_failed_or_unsupported(
    tmp_path: Path, docs: dict[str, Path]
) -> None:
    from osa.domain.errors import Unsupported

    for name, path in broken_files(tmp_path, docs).items():
        with pytest.raises((ExtractionFailed, Unsupported)) as info:
            extract(path)
        assert str(info.value), name


def test_extract_is_deterministic(
    tmp_path_factory: pytest.TempPathFactory, fixture_maker: ModuleType, docs: dict[str, Path]
) -> None:
    other = fixture_maker.build(tmp_path_factory.mktemp("docs_again"))

    for name in ("texto.pdf", "misto.pdf", "contrato.docx", "planilha.xlsx", "imagem.png"):
        first = extract(docs[name])
        second = extract(other[name])
        assert first.text == second.text, name
        assert first.method == second.method
        assert first.pages_total == second.pages_total
        assert [(i.page, i.width, i.height) for i in first.images] == [
            (i.page, i.width, i.height) for i in second.images
        ]
        assert [i.data for i in first.images] == [i.data for i in second.images]


def test_fixture_dates_are_fixed(docs: dict[str, Path]) -> None:
    from docx import Document

    created = Document(str(docs["contrato.docx"])).core_properties.created
    assert created is not None
    assert created.replace(tzinfo=None) == dt.datetime(2026, 1, 2, 3, 4, 5)
