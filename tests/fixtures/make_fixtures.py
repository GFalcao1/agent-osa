"""Generate the fictitious extraction fixtures (TASK-105), deterministically.

Usage: ``uv run python tests/fixtures/make_fixtures.py [output_dir]`` (default
``tests/fixtures/docs/``). The tests call :func:`build` into a temporary directory.

Everything is invented: companies, CNPJs (valid check digits, no real owner),
amounts and dates. PDFs use ReportLab's built-in Helvetica (no system font),
scans are rendered with pypdfium2, and metadata dates are fixed, so the content
(not the container bytes) is identical on every run and machine.
"""

from __future__ import annotations

import datetime as dt
import io
import sys
from pathlib import Path

import pypdfium2 as pdfium
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.lib.pdfencrypt import StandardEncryption
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

FIXED = dt.datetime(2026, 1, 2, 3, 4, 5)


def _cnpj(base12: str) -> str:
    def digit(digits: str, weights: list[int]) -> str:
        rest = sum(int(d) * w for d, w in zip(digits, weights, strict=True)) % 11
        return "0" if rest < 2 else str(11 - rest)

    first = digit(base12, [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])
    second = digit(base12 + first, [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])
    return base12 + first + second


def _mask(c: str) -> str:
    return f"{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}"


CNPJ_NOTA = _cnpj("112223330001")
CNPJ_CONTRATO = _cnpj("554443330001")
CNPJ_NOTA_MASKED = _mask(CNPJ_NOTA)
CNPJ_CONTRATO_MASKED = _mask(CNPJ_CONTRATO)
VALOR_NOTA = "12.345,67"

NOTA_LINES = [
    "NOTA FISCAL DE SERVIÇO Nº 000123",
    "Prestador: Exemplo Comércio e Serviços Ltda",
    f"CNPJ: {CNPJ_NOTA_MASKED}",
    "Tomador: Fictícia Indústria de Alimentos S.A.",
    "Descrição: manutenção preventiva de equipamentos e contratação de consultoria",
    f"Valor total: R$ {VALOR_NOTA}",
    "Vencimento: 15/11/2026",
    "Município: São José dos Pinhais, Paraná",
]
# Words the OCR and text tests look for (accents included).
NOTA_WORDS = ["NOTA", "FISCAL", "SERVIÇO", "Prestador", "Comércio", "Tomador", "manutenção"]

CONTRATO_PARAGRAPHS = [
    "CONTRATO DE PRESTAÇÃO DE SERVIÇOS",
    f"Contratante: Modelo Transportes e Logística Ltda, CNPJ {CNPJ_CONTRATO_MASKED}.",
    "Cláusula primeira: o objeto deste contrato é a locação de veículos utilitários.",
]
CONTRATO_TABLE = [["Item", "Valor mensal"], ["Locação de utilitários", "R$ 8.900,00"]]
CONTRATO_FOOTER = "RODAPE-FICTICIO-NAO-EXTRAIDO"

# A scanner stamp: a small text layer (well over min_chars_per_page) on a full-page scan.
SCANNER_STAMP = "Digitalizado em 20/10/2026 10:15 por Scanner Modelo X, lote 000123"

SEVEN_PAGES = [
    "Página {n} de 7. Anexo fictício do contrato de locação.",
    "Esta página descreve a cláusula {n}: pagamento, vencimento e rescisão.",
]


def _text_pdf(lines: list[str], *, password: str | None = None) -> bytes:
    buffer = io.BytesIO()
    encrypt = StandardEncryption(password, canPrint=1) if password else None
    pdf = canvas.Canvas(buffer, pagesize=A4, invariant=1, encrypt=encrypt)
    pdf.setTitle("Fixture fictícia")
    _draw_page(pdf, lines)
    pdf.save()
    return buffer.getvalue()


def _draw_page(
    pdf: canvas.Canvas, lines: list[str], font: str = "Helvetica", size: int = 12
) -> None:
    pdf.setFont(font, size)
    y = 800.0
    for line in lines:
        pdf.drawString(56, y, line)
        y -= size * 2.2
    pdf.showPage()


def _render_png(lines: list[str], dpi: int = 200) -> bytes:
    """A4 page of text as a grayscale PNG, as a scanner would produce it."""
    document = pdfium.PdfDocument(_text_pdf(lines))
    try:
        page = document[0]
        bitmap = page.render(scale=dpi / 72)
        image = bitmap.to_pil().convert("L")
        page.close()
    finally:
        document.close()
    buffer = io.BytesIO()
    image.save(buffer, format="PNG", optimize=True)
    return buffer.getvalue()


def _image_only_pdf(pages: list[bytes]) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4, invariant=1)
    for png in pages:
        pdf.drawImage(ImageReader(io.BytesIO(png)), 0, 0, width=A4[0], height=A4[1])
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def _stamped_scan_pdf(png: bytes) -> bytes:
    """A scanned page whose only text layer is the scanner's stamp in the margin."""
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4, invariant=1)
    pdf.drawImage(ImageReader(io.BytesIO(png)), 0, 0, width=A4[0], height=A4[1])
    pdf.setFont("Helvetica", 6)
    pdf.drawString(20, 10, SCANNER_STAMP)
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def _mixed_pdf() -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4, invariant=1)
    _draw_page(pdf, ["Página 1 com camada de texto.", "Documento misto fictício, parte textual."])
    scan = _render_png(
        ["Página 2 somente imagem.", f"CNPJ {CNPJ_CONTRATO_MASKED} na digitalização."]
    )
    pdf.drawImage(ImageReader(io.BytesIO(scan)), 0, 0, width=A4[0], height=A4[1])
    pdf.showPage()
    _draw_page(pdf, ["Página 3 com camada de texto.", "Encerramento do documento misto fictício."])
    pdf.save()
    return buffer.getvalue()


def _docx() -> bytes:
    from docx import Document

    document = Document()
    props = document.core_properties
    props.created = props.modified = props.last_printed = FIXED
    props.author = "Autor Fictício"
    for paragraph in CONTRATO_PARAGRAPHS:
        document.add_paragraph(paragraph)
    table = document.add_table(rows=len(CONTRATO_TABLE) + 1, cols=2)
    for r, row in enumerate(CONTRATO_TABLE):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    merged = table.cell(2, 0).merge(table.cell(2, 1))
    merged.text = "Observação: reajuste anual pelo índice contratual."
    document.add_paragraph("Assinado em Curitiba, em 02/03/2026.")
    document.sections[0].footer.paragraphs[0].text = CONTRATO_FOOTER
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.properties.created = workbook.properties.modified = FIXED
    first = workbook.active
    assert first is not None
    first.title = "Pagamentos"
    first["A1"], first["B1"], first["C1"] = "Fornecedor", "Valor", "Data"
    first["A2"], first["B2"], first["C2"] = "Exemplo Peças Ltda", 12345.67, dt.date(2026, 10, 5)
    second = workbook.create_sheet("Notas")
    second["A1"] = "Observação"
    second["A2"] = "Conferido pelo setor financeiro"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _heic(png: bytes) -> bytes | None:
    """HEIC needs an encoder; the dev-only pillow-heif wheels carry one."""
    try:
        from pillow_heif import register_heif_opener
    except ImportError:  # pragma: no cover
        return None
    register_heif_opener()
    buffer = io.BytesIO()
    Image.open(io.BytesIO(png)).convert("RGB").save(buffer, format="HEIF", quality=90)
    return buffer.getvalue()


def _gradient() -> bytes:
    image = Image.linear_gradient("L").resize((320, 240)).convert("RGB")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def build(out_dir: Path) -> dict[str, Path]:
    """Write every fixture into ``out_dir`` and return ``{name: path}``."""
    out_dir.mkdir(parents=True, exist_ok=True)
    nota_png = _render_png(NOTA_LINES)
    seven = [_render_png([line.format(n=n) for line in SEVEN_PAGES], dpi=200) for n in range(1, 8)]
    jpeg = io.BytesIO()
    Image.open(io.BytesIO(nota_png)).convert("RGB").save(jpeg, format="JPEG", quality=85)
    files: dict[str, bytes | None] = {
        "texto.pdf": _text_pdf(NOTA_LINES),
        "protegido.pdf": _text_pdf(NOTA_LINES, password="senha-ficticia"),
        "escaneado.pdf": _image_only_pdf([nota_png]),
        "escaneado_7p.pdf": _image_only_pdf(seven),
        "escaneado_carimbo.pdf": _stamped_scan_pdf(nota_png),
        "misto.pdf": _mixed_pdf(),
        "contrato.docx": _docx(),
        "planilha.xlsx": _xlsx(),
        "imagem.png": nota_png,
        "imagem.jpg": jpeg.getvalue(),
        "imagem.heic": _heic(nota_png),
        "gradiente.png": _gradient(),
        "certificado.pfx": b"",
        "arquivo.xyz": b"conteudo ficticio sem formato conhecido\n",
    }
    written: dict[str, Path] = {}
    for name, data in files.items():
        if data is None:
            continue
        path = out_dir / name
        path.write_bytes(data)
        written[name] = path
    return written


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent / "docs"
    for fixture_name, fixture_path in build(target).items():
        print(f"{fixture_path} ({fixture_path.stat().st_size} bytes)")
