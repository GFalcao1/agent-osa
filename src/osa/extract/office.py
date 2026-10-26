"""DOCX and XLSX extraction (python-docx, openpyxl) behind the OOXML guards.

Known limits: DOCX headers, footers, text boxes and notes are not extracted;
XLSX formulas yield their cached value (empty when the file never calculated it).
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Any

from osa.domain.errors import ExtractionFailed
from osa.extract import zipguard
from osa.extract.model import ExtractedContent, ExtractLimits, TextBudget


def _guard(path: Path, limits: ExtractLimits) -> None:
    """Run the entry-count, declared-size, effective-read and safe-XML checks."""
    with path.open("rb") as handle:
        if handle.read(8).startswith(zipguard.OLE_MAGIC):
            raise ExtractionFailed("arquivo protegido por senha ou em formato antigo")
    archive = zipguard.open_checked(path, limits)
    try:
        zipguard.scan_archive(archive, limits)
    finally:
        archive.close()


def extract_docx(path: Path, limits: ExtractLimits) -> ExtractedContent:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _guard(path, limits)
    budget = TextBudget(limits.max_text_chars)
    try:
        document = Document(str(path))
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                budget.add(block.text.strip())
            elif isinstance(block, Table):
                # A merged cell is one <w:tc> repeated by ``row.cells``: across columns
                # it is skipped, in the rows below it becomes an empty column. The
                # elements are kept (not their id()), so identity stays valid.
                above: set[Any] = set()
                for row in block.rows:
                    cells: list[str] = []
                    in_row: set[Any] = set()
                    for cell in row.cells:
                        tc = cell._tc
                        if tc in in_row:
                            continue
                        in_row.add(tc)
                        cells.append("" if tc in above else cell.text.strip())
                    above |= in_row
                    budget.add(" | ".join(cells))
                    if budget.truncated:
                        break
            if budget.truncated:
                break
    except (ExtractionFailed, MemoryError):
        raise
    except Exception as exc:
        raise ExtractionFailed("DOCX inválido") from exc
    # pages_total=1 is a convention: DOCX has no reliable pagination without rendering.
    return ExtractedContent(
        text=budget.text(),
        images=(),
        pages_total=1,
        truncated=budget.truncated,
        method="docx",
    )


def _cell_text(value: Any) -> str:
    if isinstance(value, dt.datetime) and value.time() == dt.time():
        return value.date().isoformat()
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return str(value).strip()


def extract_xlsx(path: Path, limits: ExtractLimits) -> ExtractedContent:
    from openpyxl import load_workbook

    _guard(path, limits)
    budget = TextBudget(limits.max_text_chars)
    truncated = False
    try:
        workbook = load_workbook(str(path), read_only=True, data_only=True)
    except MemoryError:
        raise
    except Exception as exc:
        raise ExtractionFailed("XLSX inválido") from exc
    try:
        sheets = list(workbook.worksheets)
        # Every tab counts, chartsheets included; a workbook always has at least one.
        pages_total = max(1, len(workbook.sheetnames))
        if len(sheets) > limits.max_xlsx_sheets:
            truncated = True
        for sheet in sheets[: limits.max_xlsx_sheets]:
            rows = sheet.iter_rows(max_col=limits.max_xlsx_cols)
            for number, row in enumerate(rows, start=1):
                if number > limits.max_xlsx_rows:
                    truncated = True
                    break
                for cell in row:
                    value = getattr(cell, "value", None)
                    if value is None or value == "":
                        continue
                    budget.add(f"{sheet.title}!{cell.coordinate}: {_cell_text(value)}")
                    if budget.truncated:
                        break
                if budget.truncated:
                    break
            if budget.truncated:
                break
    except (ExtractionFailed, MemoryError):
        raise
    except Exception as exc:
        raise ExtractionFailed("XLSX inválido") from exc
    finally:
        workbook.close()
    return ExtractedContent(
        text=budget.text(),
        images=(),
        pages_total=pages_total,
        truncated=truncated or budget.truncated,
        method="xlsx",
    )
